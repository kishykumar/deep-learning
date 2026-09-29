"""Run the MNIST MLP ablations from the CMU 11-785 Lab 1 (Part 4) sample sheet.
Generated alongside simple_classification_ablation.ipynb (same code). See the notebook for full instructions.

Usage:  python run_ablations.py [--row 5 | --member 2 | --custom --lr ... ] [--force]
"""
# %% Setup: imports and lookup tables
import os, json, time, math, random
from datetime import datetime

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.nn.functional as F
from torchvision import datasets

INPUT_SIZE, OUTPUT_SIZE = 28 * 28, 10

# Hidden-layer widths for each architecture (base width 128, as in the CMU lab).
# Edit or add entries here to define your own shapes.
ARCHITECTURES = {
    "Simple":           [128],
    "Cylinder":         [128, 128],
    "Pyramid":          [128, 64],
    "Inverted Pyramid": [64, 128],
    "Cone":             [256, 128, 64, 32],
    "Hourglass":        [128, 32, 128],
}
ACTIVATIONS = {"ReLU": nn.ReLU, "Sigmoid": nn.Sigmoid, "Tanh": nn.Tanh}
OPTIMIZERS = ["SGD", "Adam", "AdamW", "RMSprop"]

DEFAULT_CFG = dict(lr=1e-3, batch_size=64, epochs=3, weight_decay=0.0, activation="ReLU",
                   optimizer="Adam", architecture="Pyramid", momentum=0.0, seed=0)


def _canon(name, options, alias=None):
    """Case-insensitive lookup of a valid option name."""
    key = str(name).strip().lower().replace("_", " ").replace("-", " ")
    key = (alias or {}).get(key, key)
    for opt in options:
        if opt.lower() == key:
            return opt
    raise ValueError(f"Unknown value {name!r}. Valid options: {list(options)}")


# %% Data: MNIST held in memory on the compute device
def get_device(pref="auto"):
    if pref != "auto":
        return torch.device(pref)
    if torch.cuda.is_available():
        return torch.device("cuda")
    if torch.backends.mps.is_available():
        return torch.device("mps")
    return torch.device("cpu")


_DATA_CACHE = {}

def load_mnist(data_root, normalize, device):
    """Returns (x_train, y_train, x_val, y_val). Uses the local MNIST files (no download).
    'Validation' = the 10,000-image MNIST test split, as in the CMU lab."""
    key = (os.path.abspath(data_root), normalize, str(device))
    if key not in _DATA_CACHE:
        out = []
        for train in (True, False):
            ds = datasets.MNIST(root=data_root, train=train, download=False)
            x = ds.data.float().div(255.0).unsqueeze(1)          # [N, 1, 28, 28]
            if normalize:
                x = (x - 0.1307) / 0.3081
            out += [x.to(device), ds.targets.long().to(device)]
        _DATA_CACHE[key] = tuple(out)
    return _DATA_CACHE[key]


# %% Model and optimizer builders
def build_model(architecture, activation):
    widths = ARCHITECTURES[_canon(architecture, ARCHITECTURES, {"cylindar": "cylinder"})]
    act = ACTIVATIONS[_canon(activation, ACTIVATIONS)]
    layers, prev = [nn.Flatten()], INPUT_SIZE
    for w in widths:
        layers += [nn.Linear(prev, w), act()]
        prev = w
    layers.append(nn.Linear(prev, OUTPUT_SIZE))
    return nn.Sequential(*layers)


def build_optimizer(name, params, lr, weight_decay, momentum=0.0):
    name = _canon(name, OPTIMIZERS)
    if name == "SGD":
        return torch.optim.SGD(params, lr=lr, weight_decay=weight_decay, momentum=momentum)
    if name == "Adam":
        return torch.optim.Adam(params, lr=lr, weight_decay=weight_decay)
    if name == "AdamW":
        return torch.optim.AdamW(params, lr=lr, weight_decay=weight_decay)
    return torch.optim.RMSprop(params, lr=lr, weight_decay=weight_decay)


# %% Training and evaluation
def train_one_epoch(model, x, y, batch_size, optimizer, gen):
    """One shuffled pass over the training set. Returns (mean loss, accuracy) over the epoch."""
    model.train()
    n = x.size(0)
    perm = torch.randperm(n, generator=gen).to(x.device)
    loss_sum = torch.zeros((), device=x.device)
    correct = torch.zeros((), device=x.device)
    for s in range(0, n, batch_size):
        idx = perm[s:s + batch_size]
        xb, yb = x[idx], y[idx]
        logits = model(xb)
        loss = F.cross_entropy(logits, yb)
        optimizer.zero_grad(set_to_none=True)
        loss.backward()
        optimizer.step()
        loss_sum += loss.detach() * xb.size(0)
        correct += (logits.argmax(1) == yb).sum().float()
    return loss_sum.item() / n, correct.item() / n


@torch.no_grad()
def evaluate(model, x, y, chunk=5000):
    model.eval()
    n = x.size(0)
    loss_sum = torch.zeros((), device=x.device)
    correct = torch.zeros((), device=x.device)
    for s in range(0, n, chunk):
        xb, yb = x[s:s + chunk], y[s:s + chunk]
        logits = model(xb)
        loss_sum += F.cross_entropy(logits, yb) * xb.size(0)
        correct += (logits.argmax(1) == yb).sum().float()
    return loss_sum.item() / n, correct.item() / n


# %% run_experiment: train one configuration end to end
def run_experiment(cfg, data_root="data", device_pref="auto", normalize=True, verbose=True):
    """Train one config. Returns final-epoch metrics, the per-epoch history and auto-notes."""
    cfg = {**DEFAULT_CFG, **cfg}
    device = get_device(device_pref)
    x_tr, y_tr, x_va, y_va = load_mnist(data_root, normalize, device)

    seed = int(cfg["seed"])
    random.seed(seed); np.random.seed(seed); torch.manual_seed(seed)
    gen = torch.Generator().manual_seed(seed)          # controls the shuffling order

    model = build_model(cfg["architecture"], cfg["activation"]).to(device)
    opt = build_optimizer(cfg["optimizer"], model.parameters(), cfg["lr"],
                          cfg["weight_decay"], cfg["momentum"])

    history, diverged_at, t0 = [], None, time.time()
    for ep in range(1, int(cfg["epochs"]) + 1):
        tl, ta = train_one_epoch(model, x_tr, y_tr, int(cfg["batch_size"]), opt, gen)
        vl, va = evaluate(model, x_va, y_va)
        history.append(dict(epoch=ep, train_loss=tl, train_acc=ta, val_loss=vl, val_acc=va))
        if verbose:
            print(f"  epoch {ep:>2}/{cfg['epochs']}: train_loss={tl:.4f} train_acc={ta:.4f} "
                  f"| val_loss={vl:.4f} val_acc={va:.4f}")
        if not (math.isfinite(tl) and math.isfinite(vl)):
            diverged_at = ep
            break

    last = history[-1]                                   # FINAL epoch is what gets reported
    best = max(history, key=lambda h: h["val_acc"])
    notes = [f"best val acc {best['val_acc']:.4f} at epoch {best['epoch']}"]
    if diverged_at:
        notes.append(f"loss became NaN/inf at epoch {diverged_at} (stopped early)")
    elif last["val_acc"] < 0.5:
        notes.append("did not converge (final val acc < 50%)")
    return dict(train_loss=last["train_loss"], train_acc=last["train_acc"],
                val_loss=last["val_loss"], val_acc=last["val_acc"],
                best_val_acc=best["val_acc"], best_epoch=best["epoch"],
                seconds=time.time() - t0, device=str(device),
                notes="; ".join(notes), history=history)


# %% The 12 ablation rows from the CMU lab sample sheet (Part 4)
# (member, learning rate, batch size, epochs, regularization (weight decay), activation, optimizer, architecture)
_SHEET = [
    (1, 0.001, 32,  5, 0.0,  "ReLU",    "SGD",     "Pyramid"),
    (1, 0.001, 64,  5, 0.01, "ReLU",    "SGD",     "Pyramid"),
    (1, 0.01,  32, 10, 0.0,  "Sigmoid", "SGD",     "Inverted Pyramid"),
    (2, 0.01,  64, 10, 0.01, "Sigmoid", "SGD",     "Inverted Pyramid"),
    (2, 0.1,  128, 20, 0.0,  "Tanh",    "Adam",    "Cone"),
    (2, 0.1,  128, 20, 0.1,  "Tanh",    "Adam",    "Cone"),
    (3, 0.001, 32, 20, 0.1,  "ReLU",    "Adam",    "Hourglass"),
    (3, 0.01, 128,  5, 0.0,  "Tanh",    "RMSprop", "Hourglass"),
    (3, 0.01, 128,  5, 0.01, "Sigmoid", "RMSprop", "Cylinder"),
    (4, 0.001, 128, 20, 0.1, "ReLU",    "Adam",    "Cylinder"),
    (4, 0.1,   32,  5, 0.0,  "Sigmoid", "Adam",    "Pyramid"),
    (4, 0.1,   64, 10, 0.01, "Tanh",    "SGD",     "Pyramid"),
]
ROWS = [dict(row=i + 1, member=m, lr=lr, batch_size=bs, epochs=ep, weight_decay=wd,
             activation=act, optimizer=opt, architecture=arch)
        for i, (m, lr, bs, ep, wd, act, opt, arch) in enumerate(_SHEET)]


def select_rows(select):
    """select: 'all' | 'row:5' | 'rows:1,2,3' | 'member:2'"""
    s = str(select).strip().lower()
    if s == "all":
        return list(ROWS)
    kind, _, val = s.partition(":")
    ids = [int(t) for t in val.replace(" ", "").split(",") if t]
    if kind in ("row", "rows"):
        picked = [r for r in ROWS if r["row"] in ids]
    elif kind in ("member", "members"):
        picked = [r for r in ROWS if r["member"] in ids]
    else:
        raise ValueError("SELECT must be 'all', 'row:5', 'rows:1,2,3', 'member:2' or 'custom'")
    if not picked:
        raise ValueError(f"No ablation rows match {select!r} (rows are numbered 1-{len(ROWS)}).")
    return picked


# %% Saving results: JSON store (for resume), CSV/TSV, and the .xlsx sheet
SHEET_HEADERS = ["Ablations for...", "Learning Rate", "Batch Size", "Epochs", "Regularization Strength",
                 "Activation Function", "Optimizer", "Architecture", "Training Loss", "Validation Loss",
                 "Accuracy", "Notes", "Row #", "Training Accuracy", "Best Val Accuracy", "Best Val Epoch",
                 "Time (s)", "Device", "Seed"]
FLAT_COLUMNS = ["Row #", "Team Member", "Learning Rate", "Batch Size", "Epochs", "Regularization Strength",
                "Activation Function", "Optimizer", "Architecture", "Training Loss", "Validation Loss",
                "Accuracy", "Notes", "Training Accuracy", "Best Val Accuracy", "Best Val Epoch",
                "Time (s)", "Device", "Seed"]


def make_record(row_id, member, cfg, res, seed):
    return {"Row #": row_id, "Team Member": member, "Learning Rate": cfg["lr"], "Batch Size": cfg["batch_size"],
            "Epochs": cfg["epochs"], "Regularization Strength": cfg["weight_decay"],
            "Activation Function": cfg["activation"], "Optimizer": cfg["optimizer"],
            "Architecture": cfg["architecture"], "Training Loss": res["train_loss"],
            "Validation Loss": res["val_loss"], "Accuracy": res["val_acc"], "Notes": res["notes"],
            "Training Accuracy": res["train_acc"], "Best Val Accuracy": res["best_val_acc"],
            "Best Val Epoch": res["best_epoch"], "Time (s)": round(res["seconds"], 1),
            "Device": res["device"], "Seed": seed}


def load_state(out_dir):
    p = os.path.join(out_dir, "results.json")
    if os.path.exists(p):
        with open(p) as f:
            return json.load(f)
    return {"records": {}, "history": {}}


def save_state(state, out_dir):
    with open(os.path.join(out_dir, "results.json"), "w") as f:
        json.dump(state, f, indent=1)


def _num(v):
    """Excel can't store NaN/inf, so write those as text."""
    if isinstance(v, float) and not math.isfinite(v):
        return "NaN" if math.isnan(v) else "inf"
    return v


def write_outputs(state, out_dir):
    recs = [state["records"][k] for k in sorted(state["records"], key=int)]
    df = pd.DataFrame(recs, columns=FLAT_COLUMNS)
    df.to_csv(os.path.join(out_dir, "ablation_results.csv"), index=False)
    df.to_csv(os.path.join(out_dir, "ablation_results.tsv"), index=False, sep="\t")
    try:
        write_xlsx(state, os.path.join(out_dir, "ablation_results.xlsx"))
    except PermissionError:
        print("  (could not write the .xlsx - is it open in Excel? Close it and re-run.)")


def write_xlsx(state, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Ablations"
    head_fill = PatternFill("solid", fgColor="FFD966")
    band_fill = PatternFill("solid", fgColor="FFF2CC")
    side = Side(style="thin", color="999999")
    border = Border(left=side, right=side, top=side, bottom=side)

    for c, h in enumerate(SHEET_HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.border = Font(bold=True), head_fill, border
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    pending = "to be filled"
    for i, row in enumerate(ROWS):
        r = i + 2
        rec = state["records"].get(str(row["row"]))
        vals = [f"Team Member {row['member']}", row["lr"], row["batch_size"], row["epochs"], row["weight_decay"],
                row["activation"], row["optimizer"], row["architecture"]]
        if rec:
            vals += [rec["Training Loss"], rec["Validation Loss"], rec["Accuracy"], rec["Notes"], row["row"],
                     rec["Training Accuracy"], rec["Best Val Accuracy"], rec["Best Val Epoch"],
                     rec["Time (s)"], rec["Device"], rec["Seed"]]
        else:
            vals += [pending] * 4 + [row["row"]] + [pending] * 6
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=_num(v))
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=(c == 12))
            if row["member"] % 2 == 0:
                cell.fill = band_fill
            if c in (9, 10):
                cell.number_format = "0.0000"
            if c in (11, 14, 15):
                cell.number_format = "0.00%"
            if v == pending:
                cell.font = Font(color="999999")

    r0 = 2                                   # merge the "Team Member N" cell across each member's rows
    while r0 < len(ROWS) + 2:
        m = ROWS[r0 - 2]["member"]
        r1 = r0
        while r1 + 1 < len(ROWS) + 2 and ROWS[r1 - 1]["member"] == m:
            r1 += 1
        ws.merge_cells(start_row=r0, start_column=1, end_row=r1, end_column=1)
        r0 = r1 + 1

    for c, w in enumerate([16, 10, 10, 8, 14, 12, 11, 17, 12, 12, 11, 48, 7, 12, 12, 10, 9, 8, 6], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "B2"

    hs = wb.create_sheet("Per-epoch history")
    hs.append(["Row #", "Team Member", "Epoch", "Train Loss", "Train Acc", "Val Loss", "Val Acc"])
    for c in hs[1]:
        c.font, c.fill = Font(bold=True), head_fill
    for k in sorted(state["history"], key=int):
        member = ROWS[int(k) - 1]["member"]
        for h in state["history"][k]:
            hs.append([int(k), member, h["epoch"], _num(h["train_loss"]), _num(h["train_acc"]),
                       _num(h["val_loss"]), _num(h["val_acc"])])
    hs.freeze_panes = "A2"
    wb.save(path)


def results_dataframe(out_dir="ablation_results"):
    """Load the saved results as a DataFrame (for display in the notebook)."""
    return pd.read_csv(os.path.join(out_dir, "ablation_results.csv"))


# %% run_selection: run the chosen ablation rows (or one custom config) and save everything
def run_selection(select="all", custom=None, force=False, data_root="data", out_dir="ablation_results",
                  device_pref="auto", normalize=True, seed=0, overrides=None):
    """overrides (e.g. {"epochs": 1}) is for quick smoke tests: results go to <out_dir>_override, not your sheet."""
    if overrides:
        out_dir = out_dir + "_override"
    os.makedirs(out_dir, exist_ok=True)
    print(f"device: {get_device(device_pref)} | data: {os.path.abspath(data_root)} | output: {os.path.abspath(out_dir)}")

    if str(select).strip().lower() == "custom":
        cfg = {**DEFAULT_CFG, **(custom or {}), "seed": seed, **(overrides or {})}
        print(f"Custom run: {cfg}")
        res = run_experiment(cfg, data_root, device_pref, normalize)
        rec = make_record("custom", "-", cfg, res, seed)
        rec["Timestamp"] = datetime.now().isoformat(timespec="seconds")
        p = os.path.join(out_dir, "custom_runs.csv")
        pd.DataFrame([rec]).to_csv(p, mode="a", header=not os.path.exists(p), index=False)
        print(f"FINAL  train_loss={res['train_loss']:.4f}  val_loss={res['val_loss']:.4f}  "
              f"val_acc={res['val_acc']:.4f}  | {res['notes']}\nAppended to {p}")
        return rec

    state = load_state(out_dir)
    todo = select_rows(select)
    for n, row in enumerate(todo, 1):
        key = str(row["row"])
        if key in state["records"] and not force:
            print(f"[{n}/{len(todo)}] Row {row['row']} already done - skipping (set FORCE_RERUN / --force to redo)")
            continue
        cfg = {k: row[k] for k in ("lr", "batch_size", "epochs", "weight_decay", "activation",
                                   "optimizer", "architecture")}
        cfg.update(momentum=0.0, seed=seed, **(overrides or {}))
        print(f"[{n}/{len(todo)}] Row {row['row']} (Team Member {row['member']}): {cfg}")
        res = run_experiment(cfg, data_root, device_pref, normalize)
        state["records"][key] = make_record(row["row"], row["member"], cfg, res, seed)
        state["history"][key] = res["history"]
        save_state(state, out_dir)
        write_outputs(state, out_dir)
        print(f"  -> FINAL train_loss={res['train_loss']:.4f}  val_loss={res['val_loss']:.4f}  "
              f"val_acc={res['val_acc']:.4f}  ({res['seconds']:.0f}s)")
    print(f"\nDone. Results: {os.path.abspath(out_dir)}/ablation_results.xlsx (+ .csv / .tsv)")
    return state


# ----------------------------------------------------------------------------------------
# Command-line interface
# ----------------------------------------------------------------------------------------
if __name__ == "__main__":
    import argparse
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="Run the MNIST MLP ablations from the CMU lab sample sheet.")
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--row", help="run row(s) 1-12, e.g. 5 or 1,4,9")
    g.add_argument("--member", help="run all rows of team member(s), e.g. 2 or 1,3")
    g.add_argument("--custom", action="store_true", help="run one custom config (use the options below)")
    ap.add_argument("--force", action="store_true", help="redo rows that already have results")
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--no-normalize", action="store_true")
    ap.add_argument("--device", default="auto", help="auto | cpu | mps | cuda")
    ap.add_argument("--data-root", default=os.path.join(here, "data"))
    ap.add_argument("--out-dir", default=os.path.join(here, "ablation_results"))
    ap.add_argument("--epochs-override", type=int, help="smoke test: force this many epochs (results go to <out-dir>_override)")
    d = DEFAULT_CFG
    ap.add_argument("--lr", type=float, default=d["lr"])
    ap.add_argument("--batch-size", type=int, default=d["batch_size"])
    ap.add_argument("--epochs", type=int, default=d["epochs"])
    ap.add_argument("--weight-decay", type=float, default=d["weight_decay"])
    ap.add_argument("--activation", default=d["activation"])
    ap.add_argument("--optimizer", default=d["optimizer"])
    ap.add_argument("--architecture", default=d["architecture"])
    ap.add_argument("--momentum", type=float, default=d["momentum"])
    a = ap.parse_args()

    select = "custom" if a.custom else (f"row:{a.row}" if a.row else (f"member:{a.member}" if a.member else "all"))
    custom = dict(lr=a.lr, batch_size=a.batch_size, epochs=a.epochs, weight_decay=a.weight_decay,
                  activation=a.activation, optimizer=a.optimizer, architecture=a.architecture, momentum=a.momentum)
    run_selection(select=select, custom=custom, force=a.force, data_root=a.data_root, out_dir=a.out_dir,
                  device_pref=a.device, normalize=not a.no_normalize, seed=a.seed,
                  overrides={"epochs": a.epochs_override} if a.epochs_override else None)
